import os
import shutil
import json
import time
import unittest

TEST_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'test_kafel_isolated.db')
PROD_DB = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'kafel_database.db')
os.environ["DB_FILE"] = TEST_DB

# Copy real database structure and seed data to isolated test database
if os.path.exists(PROD_DB):
    shutil.copy2(PROD_DB, TEST_DB)

import app
import database

class TestKafelAuthAndRoles(unittest.TestCase):
    @classmethod
    def tearDownClass(cls):
        # Clean up isolated test db files after tests complete
        for ext in ['', '-wal', '-shm']:
            f = TEST_DB + ext
            if os.path.exists(f):
                try:
                    os.remove(f)
                except Exception:
                    pass

    def setUp(self):
        self.client = app.app.test_client()
        app.app.config["TESTING"] = True
        app.reset_all_rate_limits()
        database.init_db()

    def test_unauthenticated_redirect(self):
        # Without logging in, visiting dashboard redirects to login
        r = self.client.get("/", follow_redirects=False)
        self.assertEqual(r.status_code, 302)
        self.assertIn("/login", r.headers["Location"])

    def test_admin_flow(self):
        # 1. Login as admin
        r = self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Boshqaruv Paneli", r.data)

        # 2. Access Xodimlar page
        r = self.client.get("/admin/xodimlar")
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Xodimlar va Rollar Boshqaruvi", r.data)

        # 3. Create a new seller employee
        r = self.client.post("/admin/xodimlar", data={
            "full_name": "Vali Sotuvchi",
            "username": "vali_sotuv",
            "password": "vali1234password",
            "role": "sotuvchi"
        }, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"vali_sotuv", r.data)

        # 4. Logout
        r = self.client.get("/logout", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Tizimga Kirish", r.data)

    def test_sotuvchi_permissions(self):
        # 1. Login as sotuvchi
        r = self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"}, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Sotuv Bo'limi", r.data)

        # 2. Sotuvchi can access Kassa
        r = self.client.get("/kassa")
        self.assertEqual(r.status_code, 200)

        # 3. Sotuvchi CANNOT access Admin Xodimlar page (redirected with warning)
        r = self.client.get("/admin/xodimlar", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r.data)

        # 4. Sotuvchi CANNOT access New Product page
        r = self.client.get("/ombor/yangi", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r.data)

        # 5. Sotuvchi CANNOT access Warehouse Inventory (/ombor)
        r_ombor = self.client.get("/ombor", follow_redirects=True)
        self.assertEqual(r_ombor.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r_ombor.data)

        # 6. Sotuvchi CANNOT access Customers & Debt CRM (/mijozlar)
        r_mijoz = self.client.get("/mijozlar", follow_redirects=True)
        self.assertEqual(r_mijoz.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r_mijoz.data)

    def test_omborchi_permissions(self):
        # 1. Login as omborchi
        r = self.client.post("/login", data={"username": "omborchi", "password": "omborchi123"}, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"Ombor - Kafellar Boshqaruvi", r.data)

        # 2. Omborchi can access Ombor yangi
        r = self.client.get("/ombor/yangi")
        self.assertEqual(r.status_code, 200)

        # 3. Omborchi can access Nakladnoylar
        r = self.client.get("/nakladnoylar")
        self.assertEqual(r.status_code, 200)

        # 4. Omborchi CANNOT access Kassa
        r = self.client.get("/kassa", follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r.data)

        # 5. Omborchi CANNOT access KPI & Reyting
        r_kpi = self.client.get("/kpi", follow_redirects=True)
        self.assertEqual(r_kpi.status_code, 200)
        self.assertIn(b"ruxsat bermaydi", r_kpi.data)

    def test_dashboard_charts_and_data(self):
        # Admin login
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r = self.client.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn(b"salesChart", r.data)
        self.assertIn(b"paymentChart", r.data)

    def test_showroom_search_api(self):
        # Sotuvchi login
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"}, follow_redirects=True)
        r = self.client.get("/api/products/search?size=60x120")
        self.assertEqual(r.status_code, 200)
        data = json.loads(r.data)
        self.assertTrue(data["success"])

    def test_checkout_and_thermal_receipt(self):
        # Admin login
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        
        # 1. Query a product
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM products WHERE quantity_in_stock > 1 LIMIT 1")
        prod = cursor.fetchone()
        conn.close()

        if prod:
            # 2. Perform checkout with discount
            r = self.client.post("/api/checkout", json={
                "customer_name": "Test Mijoz",
                "customer_phone": "+998901112233",
                "payment_method": "naqd",
                "cashier_note": "Sinov savdo",
                "discount_amount": 5000,
                "items": [{
                    "id": prod["id"],
                    "qty": 1.0,
                    "price": prod["price"]
                }]
            })
            self.assertEqual(r.status_code, 200)
            res = json.loads(r.data)
            self.assertTrue(res["success"])
            order_id = res["order_id"]

            # 3. Check 80mm thermal receipt
            r_chek = self.client.get(f"/nakladnoy/{order_id}/chek")
            self.assertEqual(r_chek.status_code, 200)
            self.assertIn(b"KAFEL CENTER", r_chek.data)
            self.assertIn(b"80mm", r_chek.data)

    def test_audit_and_tracking_flow(self):
        # 1. Login as sotuvchi and checkout
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"})
        
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM products WHERE quantity_in_stock > 1 LIMIT 1")
        prod = cursor.fetchone()
        conn.close()

        r = self.client.post("/api/checkout", json={
            "customer_name": "Audit Tekshiruv Xaridor",
            "customer_phone": "+998909998877",
            "payment_method": "naqd",
            "cashier_note": "Audit sinovi",
            "items": [{"id": prod["id"], "qty": 1.0, "price": prod["price"]}]
        })
        self.assertEqual(r.status_code, 200)
        order_id = json.loads(r.data)["order_id"]

        # 2. Check seller info recorded in DB
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM orders WHERE id = ?", (order_id,))
        order = cursor.fetchone()
        self.assertEqual(order["seller_name"], "Sotuvchi (Kassir)")
        self.assertEqual(order["status"], "tolandi")
        self.assertIsNone(order["dispatched_by"])
        conn.close()

        # 3. Login as omborchi and dispatch
        self.client.get("/logout")
        self.client.post("/login", data={"username": "omborchi", "password": "omborchi123"})

        r_disp = self.client.post(f"/api/order/{order_id}/dispatch", json={"note": "Topshirildi"})
        self.assertEqual(r_disp.status_code, 200)

        # 4. Check dispatched_by recorded
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM orders WHERE id = ?", (order_id,))
        order_after = cursor.fetchone()
        self.assertEqual(order_after["dispatched_by"], "Bosh Omborchi")
        self.assertEqual(order_after["status"], "yuk_berildi")
        self.assertIsNotNone(order_after["dispatched_at"])
        conn.close()

        # 5. Check Hisobot & Audit page loads with 200 and includes brands and broken tiles
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})
        r_rep = self.client.get("/hisobot")
        self.assertEqual(r_rep.status_code, 200)
        html_rep = r_rep.data.decode("utf-8")
        self.assertIn("Umumiy Hisobot", html_rep)
        self.assertIn("Bosh Balans Vedomosti", html_rep)
        self.assertIn("Sotuvchi (Kassir)", html_rep)
        self.assertIn("Bosh Omborchi", html_rep)
        self.assertIn("Kafel Markalari", html_rep)
        self.assertIn("Siniq & Braklar", html_rep)
        self.assertIn("Vitra Ceramic", html_rep)
        self.assertIn("Kerama Marazzi", html_rep)

    def test_kpi_dashboard_and_performance(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Access /kpi for different periods
        for p in ["today", "week", "month", "all"]:
            r = self.client.get(f"/kpi?period={p}")
            self.assertEqual(r.status_code, 200)
            html = r.data.decode("utf-8")
            self.assertIn("Xodimlar KPI", html)
            self.assertIn("Sotuvchilar Reytingi", html)

        # 3. Test bonus rate customization
        r_custom = self.client.get("/kpi?period=month&rate=3.0")
        self.assertEqual(r_custom.status_code, 200)
        self.assertIn("3.0%", r_custom.data.decode("utf-8"))

    def test_sidebar_navigation_and_roles(self):
        # 1. Admin login -> Check sidebar structure & elements
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})
        r_admin = self.client.get("/ombor")
        self.assertEqual(r_admin.status_code, 200)
        html_admin = r_admin.data.decode("utf-8")
        self.assertIn('id="sidebar"', html_admin)
        self.assertIn('id="topHeader"', html_admin)
        self.assertIn('id="desktopSidebarToggle"', html_admin)
        self.assertIn('id="mobileSidebarToggle"', html_admin)
        self.assertIn("Asosiy Boshqaruv", html_admin)
        self.assertIn("Savdo & Kassa", html_admin)
        self.assertIn("Ombor & Logistika", html_admin)
        self.assertIn("Tizim Nazorati", html_admin)
        self.assertIn("Nakladnoylar", html_admin)
        self.assertIn("Xodimlar", html_admin)

        # 2. Sotuvchi login -> Check sidebar role filtering
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"})
        r_seller = self.client.get("/sotuv")
        self.assertEqual(r_seller.status_code, 200)
        html_seller = r_seller.data.decode("utf-8")
        self.assertIn("Savdo Operatsiyalari", html_seller)
        self.assertIn("Mening KPI & Bonus", html_seller)
        self.assertNotIn("Tizim Nazorati", html_seller)
        self.assertNotIn("/admin/xodimlar", html_seller)
        self.assertNotIn("Mijozlar & Nasiya", html_seller)
        self.assertNotIn("Ombor Qoldig'i", html_seller)

        # 3. Omborchi login -> Check sidebar role filtering
        self.client.get("/logout")
        self.client.post("/login", data={"username": "omborchi", "password": "omborchi123"})
        r_ware = self.client.get("/ombor")
        self.assertEqual(r_ware.status_code, 200)
        html_ware = r_ware.data.decode("utf-8")
        side_ware = html_ware[html_ware.find('<aside'):html_ware.find('</aside>')]
        self.assertIn("Ombor & Yuk Tashish", side_ware)
        self.assertIn("Kafel Kirim Qilish", side_ware)
        self.assertNotIn("/admin/xodimlar", side_ware)
        self.assertNotIn("KPI & Reyting", side_ware)

    def test_mijozlar_crm_and_nasiya(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Add a new customer
        r = self.client.post("/mijozlar", data={
            "full_name": "Bahrom Usta",
            "phone": "+998 90 999 88 77",
            "customer_type": "usta",
            "balance_debt": "450000",
            "notes": "Yunusobod obyekti ustasi"
        }, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Bahrom Usta", r.data.decode("utf-8"))

        # 3. Search customers API
        r_search = self.client.get("/api/customers/search?q=Bahrom")
        self.assertEqual(r_search.status_code, 200)
        data = json.loads(r_search.data)
        self.assertTrue(data["success"])
        self.assertTrue(len(data["customers"]) > 0)
        cust_id = data["customers"][0]["id"]
        self.assertEqual(float(data["customers"][0]["balance_debt"]), 450000.0)

        # 4. Partial debt repayment
        r_pay = self.client.post(f"/api/customer/{cust_id}/pay-debt", json={
            "amount": 200000,
            "payment_method": "naqd",
            "note": "Oldindan qisman to'lov"
        })
        self.assertEqual(r_pay.status_code, 200)
        pay_data = json.loads(r_pay.data)
        self.assertTrue(pay_data["success"])
        self.assertEqual(pay_data["new_debt"], 250000.0)

    def test_xarajatlar_and_overhead(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Add an expense
        r = self.client.post("/xarajatlar", data={
            "category": "transport",
            "amount": "150000",
            "description": "Labo mashinasiga yoqilgi va yetkazib berish",
            "expense_date": "2026-09-14"
        }, follow_redirects=True)
        self.assertEqual(r.status_code, 200)
        self.assertIn("Labo mashinasiga yoqilgi", r.data.decode("utf-8"))

        # 3. Non-admin (sotuvchi) cannot access /xarajatlar
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"})
        r_forbidden = self.client.get("/xarajatlar", follow_redirects=True)
        self.assertEqual(r_forbidden.status_code, 200)
        self.assertIn("ruxsat bermaydi", r_forbidden.data.decode("utf-8"))

    def test_smeta_quote_page(self):
        # Sotuvchi login
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"})
        r = self.client.get("/sotuv/smeta")
        self.assertEqual(r.status_code, 200)
        html = r.data.decode("utf-8")
        self.assertIn("Smeta & Tijoriy Taklif", html)
        self.assertIn("Tavsiya etiladigan yordamchi materiallar", html)

    def test_nasiya_checkout_and_customer_debt_sync(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Get a product to sell
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, brand, model_name, price, quantity_in_stock FROM products WHERE quantity_in_stock > 2 LIMIT 1")
        p = cursor.fetchone()
        self.assertIsNotNone(p)
        p_id = p["id"]
        old_stock = p["quantity_in_stock"]
        unit_price = p["price"]

        # 3. Create or select a customer
        cursor.execute("INSERT INTO customers (full_name, phone, customer_type, balance_debt) VALUES ('Rustam Prorab', '+998901112233', 'prorab', 0)")
        cust_id = cursor.lastrowid
        conn.commit()
        conn.close()

        # 4. Perform Nasiya Checkout: 2 m² tile
        sell_qty = 2.0
        total_order = sell_qty * unit_price
        advance_paid = 50000.0
        expected_debt = total_order - advance_paid

        r_checkout = self.client.post("/api/checkout", json={
            "customer_name": "Rustam Prorab",
            "customer_phone": "+998901112233",
            "customer_id": cust_id,
            "payment_method": "nasiya",
            "is_nasiya": 1,
            "paid_amount": advance_paid,
            "debt_amount": expected_debt,
            "cashier_note": "Obyekt topshirilgach qolgani beriladi",
            "discount_amount": 0,
            "items": [{
                "id": p_id,
                "qty": sell_qty,
                "price": unit_price,
                "box_size_m2": 1.44
            }]
        })
        self.assertEqual(r_checkout.status_code, 200)
        res = json.loads(r_checkout.data)
        self.assertTrue(res["success"])
        order_id = res["order_id"]

        # 5. Verify database state
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM orders WHERE id = ?", (order_id,))
        order = cursor.fetchone()
        self.assertEqual(order["is_nasiya"], 1)
        self.assertEqual(order["debt_amount"], expected_debt)
        self.assertEqual(order["paid_amount"], advance_paid)

        # Verify stock decreased
        cursor.execute("SELECT quantity_in_stock FROM products WHERE id = ?", (p_id,))
        new_stock = cursor.fetchone()["quantity_in_stock"]
        self.assertEqual(new_stock, old_stock - sell_qty)

        # Verify customer debt balance increased
        cursor.execute("SELECT balance_debt FROM customers WHERE id = ?", (cust_id,))
        cust_debt = cursor.fetchone()["balance_debt"]
        self.assertEqual(cust_debt, expected_debt)

        conn.close()

    def test_broken_tile_transfer_and_accounting(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Insert test product 'Qora Angola' with exactly 200.0 stock
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM order_items WHERE product_id IN (SELECT id FROM products WHERE sku = 'TEST-ANGOLA-200')")
        cursor.execute("DELETE FROM broken_tiles WHERE product_id IN (SELECT id FROM products WHERE sku = 'TEST-ANGOLA-200')")
        cursor.execute("DELETE FROM stock_history WHERE product_id IN (SELECT id FROM products WHERE sku = 'TEST-ANGOLA-200')")
        cursor.execute("DELETE FROM products WHERE sku = 'TEST-ANGOLA-200'")
        cursor.execute("""
            INSERT INTO products (
                sku, brand, model_name, size,
                cost_price, price, quantity_in_stock, unit, box_size_m2, pieces_per_box
            ) VALUES (
                'TEST-ANGOLA-200', 'Kerama', 'Qora Angola Lux', '60x120',
                120000.0, 180000.0, 200.0, 'm²', 1.44, 2
            )
        """)
        product_id = cursor.lastrowid
        conn.commit()
        conn.close()

        # 3. Report defect: 1.0 m2 broken during unloading by loader Akmal
        r_report = self.client.post(f"/api/product/{product_id}/report-defect", json={
            "quantity": 1.0,
            "reason": "Yuk tushirishda sindi (yuklovchilar)",
            "responsible_person": "Yuklovchi Akmal",
            "note": "Konteynerdan tushirish paytida quti burchagi urilgan"
        })
        self.assertEqual(r_report.status_code, 200)
        res_report = json.loads(r_report.data)
        self.assertTrue(res_report["success"])
        self.assertEqual(res_report["new_stock"], 199.0)
        defect_id = res_report["defect_id"]

        # 4. Verify stock decrement in products table (200 -> 199)
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT quantity_in_stock FROM products WHERE id = ?", (product_id,))
        p_row = cursor.fetchone()
        self.assertEqual(p_row["quantity_in_stock"], 199.0)

        # 5. Verify broken_tiles ledger entry
        cursor.execute("SELECT * FROM broken_tiles WHERE id = ?", (defect_id,))
        defect_row = cursor.fetchone()
        self.assertIsNotNone(defect_row)
        self.assertEqual(defect_row["product_id"], product_id)
        self.assertEqual(defect_row["quantity"], 1.0)
        self.assertEqual(defect_row["status"], "omborda")
        self.assertEqual(defect_row["reason"], "Yuk tushirishda sindi (yuklovchilar)")
        self.assertEqual(defect_row["responsible_person"], "Yuklovchi Akmal")
        self.assertEqual(defect_row["loss_amount"], 120000.0) # 1.0 * purchase_price

        # 6. Verify audit history in stock_history
        cursor.execute("SELECT * FROM stock_history WHERE product_id = ? AND change_type = 'brak'", (product_id,))
        history_row = cursor.fetchone()
        self.assertIsNotNone(history_row)
        self.assertEqual(history_row["quantity_change"], -1.0)

        # 7. Verify auto-recorded expense entry
        cursor.execute("SELECT * FROM expenses WHERE category = 'brak' ORDER BY id DESC LIMIT 1")
        expense_row = cursor.fetchone()
        self.assertIsNotNone(expense_row)
        self.assertEqual(expense_row["amount"], 120000.0)
        self.assertIn("Qora Angola Lux", expense_row["description"])
        conn.close()

        # 8. Verify Siniqlar page displays the broken tile
        r_page = self.client.get("/ombor/siniqlar")
        self.assertEqual(r_page.status_code, 200)
        page_html = r_page.data.decode("utf-8")
        self.assertIn("Qora Angola Lux", page_html)
        self.assertIn("Yuklovchi Akmal", page_html)
        self.assertIn("Siniqlar & Brak Kafellar", page_html)

        # 9. Test disposition action: discounted sale
        r_action = self.client.post(f"/api/defect/{defect_id}/action", json={
            "action_type": "arzonlashtirib_sotildi",
            "sale_price": 50000.0,
            "action_note": "Kesib ishlatish uchun ustaga sotildi"
        })
        self.assertEqual(r_action.status_code, 200)
        res_action = json.loads(r_action.data)
        self.assertTrue(res_action["success"])

        # 10. Verify status updated to arzonlashtirib_sotildi and loss updated
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT status, loss_amount FROM broken_tiles WHERE id = ?", (defect_id,))
        updated_defect = cursor.fetchone()
        self.assertEqual(updated_defect["status"], "arzonlashtirib_sotildi")
        # 120,000 cost - 50,000 recovered = 70,000 net loss
        self.assertEqual(updated_defect["loss_amount"], 70000.0)
        conn.close()

    def test_broken_tile_discount_price_and_pos_checkout(self):
        """Test admin setting discounted price on broken tile and regular checkout via POS without double inventory deduction"""
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)

        # 1. Create a product with 200 units
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM broken_tiles WHERE product_id IN (SELECT id FROM products WHERE sku = 'ANG-TEST-002')")
        cursor.execute("DELETE FROM order_items WHERE product_id IN (SELECT id FROM products WHERE sku = 'ANG-TEST-002')")
        cursor.execute("DELETE FROM stock_history WHERE product_id IN (SELECT id FROM products WHERE sku = 'ANG-TEST-002')")
        cursor.execute("DELETE FROM products WHERE sku = 'ANG-TEST-002'")
        cursor.execute("""
            INSERT INTO products (
                sku, brand, model_name, size,
                cost_price, price, quantity_in_stock, unit, box_size_m2, pieces_per_box
            ) VALUES (
                'ANG-TEST-002', 'Cersanit', 'Qora Angola Premium', '60x120',
                100000.0, 150000.0, 200.0, 'm²', 1.44, 2
            )
        """)
        product_id = cursor.lastrowid
        conn.commit()
        conn.close()

        # 2. Break 1 unit (200 -> 199) and set initial discounted price
        r_defect = self.client.post(f"/api/product/{product_id}/report-defect", json={
            "quantity": 1.0,
            "unit": "m2",
            "reason": "Yuk tushirishda sindi (yuklovchilar)",
            "responsible_person": "Yuklovchi Dilshod",
            "discounted_price": 50000.0,
            "note": "Burchagi singan, 50 000 somdan sotiladi"
        })
        self.assertEqual(r_defect.status_code, 200)
        def_res = json.loads(r_defect.data)
        defect_id = def_res["defect_id"]
        self.assertEqual(def_res["new_stock"], 199.0)

        # 3. Admin updates discounted price to 45,000 so'm
        r_price = self.client.post(f"/api/defect/{defect_id}/set-price", json={
            "discounted_price": 45000.0
        })
        self.assertEqual(r_price.status_code, 200)
        price_res = json.loads(r_price.data)
        self.assertTrue(price_res["success"])
        self.assertEqual(price_res["new_price"], 45000.0)

        # 4. Check that GET /api/defects/for-sale lists this defect with updated price
        r_for_sale = self.client.get("/api/defects/for-sale")
        self.assertEqual(r_for_sale.status_code, 200)
        sale_data = json.loads(r_for_sale.data)
        self.assertTrue(sale_data["success"])
        found = next((d for d in sale_data["defects"] if d["id"] == defect_id), None)
        self.assertIsNotNone(found)
        self.assertEqual(found["discounted_price"], 45000.0)
        self.assertEqual(found["original_price"], 150000.0)
        self.assertEqual(found["quantity"], 1.0)

        # 5. Sell the broken tile through standard checkout (/api/checkout)
        r_checkout = self.client.post("/api/checkout", json={
            "items": [
                {
                    "id": product_id,
                    "defect_id": defect_id,
                    "is_defect": 1,
                    "qty": 1.0,
                    "price": 45000.0
                }
            ],
            "customer_id": None,
            "customer_name": "Siniq Kafel Oluvchi Usta",
            "payment_method": "naqd",
            "discount_amount": 0,
            "amount_received": 45000.0,
            "is_nasiya": False
        })
        self.assertEqual(r_checkout.status_code, 200)
        check_res = json.loads(r_checkout.data)
        self.assertTrue(check_res["success"])
        order_id = check_res["order_id"]

        # 6. Verify broken_tiles table state
        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM broken_tiles WHERE id = ?", (defect_id,))
        def_row = cursor.fetchone()
        self.assertEqual(def_row["status"], "arzonlashtirib_sotildi")
        self.assertEqual(def_row["quantity"], 0.0)
        self.assertEqual(def_row["sold_price"], 45000.0)
        self.assertEqual(def_row["order_id"], order_id)
        # Loss was 100,000 (purchase cost) - 45,000 (sold) = 55,000
        self.assertEqual(def_row["loss_amount"], 55000.0)

        # 7. CRITICAL: Verify main products table quantity did NOT double decrement!
        # Initial: 200, broken: -1 -> 199. Sold broken tile: stock must STILL be 199!
        cursor.execute("SELECT quantity_in_stock FROM products WHERE id = ?", (product_id,))
        p_row = cursor.fetchone()
        self.assertEqual(p_row["quantity_in_stock"], 199.0)

        # 8. Verify order_items has is_defect and defect_id recorded
        cursor.execute("SELECT * FROM order_items WHERE order_id = ?", (order_id,))
        oi_row = cursor.fetchone()
        self.assertEqual(oi_row["is_defect"], 1)
        self.assertEqual(oi_row["defect_id"], defect_id)
        self.assertEqual(oi_row["unit_price"], 45000.0)
        self.assertIn("Siniq/Brak", oi_row["model_name"])

        # 9. Verify stock_history logged sotuv_brak
        cursor.execute("SELECT * FROM stock_history WHERE product_id = ? AND change_type = 'sotuv_brak'", (product_id,))
        hist_row = cursor.fetchone()
        self.assertIsNotNone(hist_row)
        self.assertEqual(hist_row["quantity_change"], -1.0)
        self.assertIn("Siniq/Brak sotuv", hist_row["note"])

        # 10. Verify defect is no longer returned in /api/defects/for-sale (since it is sold)
        r_for_sale_after = self.client.get("/api/defects/for-sale")
        sale_data_after = json.loads(r_for_sale_after.data)
        found_after = next((d for d in sale_data_after["defects"] if d["id"] == defect_id), None)
        self.assertIsNone(found_after)

        conn.close()

    def test_factories_management(self):
        """Test Zavodlar & Markalar page, API CRUD, and dynamic appearance in /ombor/yangi"""
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)

        # 1. Access /zavodlar page
        r1 = self.client.get("/zavodlar")
        self.assertEqual(r1.status_code, 200)
        self.assertIn(b"Zavodlar", r1.data)

        # 2. Add a new factory via API
        test_factory_name = f"Test Zavod {int(time.time()*1000)}"
        r2 = self.client.post("/api/factories/add", json={
            "name": test_factory_name,
            "country": "O'zbekiston",
            "contact_person": "Jahongir aka",
            "phone": "+998901239988",
            "address": "Navoiy viloyati",
            "note": "Yangi granit zavodi"
        })
        self.assertEqual(r2.status_code, 200)
        d2 = json.loads(r2.data)
        self.assertTrue(d2["success"])
        factory_id = d2["id"]

        # 3. Verify it is listed in /api/factories/list
        r3 = self.client.get("/api/factories/list")
        self.assertEqual(r3.status_code, 200)
        d3 = json.loads(r3.data)
        self.assertTrue(any(f["name"] == test_factory_name for f in d3["factories"]))

        # 4. Verify it dynamically appears on /ombor/yangi
        r4 = self.client.get("/ombor/yangi")
        self.assertEqual(r4.status_code, 200)
        self.assertIn(test_factory_name, r4.data.decode("utf-8"))

        # 5. Edit the factory
        r5 = self.client.post(f"/api/factories/{factory_id}/edit", json={
            "name": test_factory_name + " Yangilangan",
            "country": "O'zbekiston",
            "contact_person": "Jahongir aka (Direktor)",
            "phone": "+998901239988",
            "address": "Navoiy EIZ",
            "note": "Navoiy yangi zavod",
            "is_active": 1
        })
        self.assertEqual(r5.status_code, 200)
        self.assertTrue(json.loads(r5.data)["success"])

        # 6. Delete the factory
        r6 = self.client.post(f"/api/factories/{factory_id}/delete")
        self.assertEqual(r6.status_code, 200)
        self.assertTrue(json.loads(r6.data)["success"])

    def test_usta_kpi_and_ratings_multi_periods(self):
        # 1. Login as admin
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"})

        # 2. Add 2 ustalar via /api/ustalar/quick-add
        r1 = self.client.post("/api/ustalar/quick-add", json={
            "full_name": "Usta Sherzod",
            "phone": "+998901112233",
            "notes": "Premium usta"
        })
        self.assertEqual(r1.status_code, 200)
        d1 = json.loads(r1.data)
        self.assertTrue(d1["success"])
        usta1_id = d1["usta"]["id"]

        r2 = self.client.post("/api/ustalar/quick-add", json={
            "full_name": "Usta Dilshod",
            "phone": "+998904445566",
            "notes": "Katta prorab"
        })
        self.assertEqual(r2.status_code, 200)
        d2 = json.loads(r2.data)
        self.assertTrue(d2["success"])
        usta2_id = d2["usta"]["id"]

        # Verify /api/ustalar lists both
        r_list = self.client.get("/api/ustalar")
        self.assertEqual(r_list.status_code, 200)
        d_list = json.loads(r_list.data)
        self.assertTrue(any(u["id"] == usta1_id for u in d_list["ustalar"]))

        # 3. Get a product to sell
        conn = database.get_db()
        c = conn.cursor()
        c.execute("SELECT id, brand, model_name, price, quantity_in_stock FROM products WHERE quantity_in_stock >= 10 LIMIT 1")
        prod = c.fetchone()
        conn.close()

        if prod:
            p_id = prod["id"]
            p_price = prod["price"]

            # Usta 1 brings 2 different clients
            self.client.post("/api/checkout", json={
                "customer_name": "Mijoz Jasur",
                "customer_phone": "+998931110001",
                "usta_id": usta1_id,
                "payment_method": "naqd",
                "items": [{"id": p_id, "brand": prod["brand"], "model_name": prod["model_name"], "price": p_price, "qty": 1.0}]
            })
            self.client.post("/api/checkout", json={
                "customer_name": "Mijoz Botir",
                "customer_phone": "+998931110002",
                "usta_id": usta1_id,
                "payment_method": "naqd",
                "items": [{"id": p_id, "brand": prod["brand"], "model_name": prod["model_name"], "price": p_price, "qty": 1.0}]
            })

            # Usta 2 brings 1 client with larger order
            r_ord = self.client.post("/api/checkout", json={
                "customer_name": "Mijoz Umid",
                "customer_phone": "+998931110003",
                "usta_id": usta2_id,
                "payment_method": "naqd",
                "items": [{"id": p_id, "brand": prod["brand"], "model_name": prod["model_name"], "price": p_price, "qty": 3.0}]
            })
            ord_id = json.loads(r_ord.data).get("order_id")

            # Verify nakladnoy has usta info
            if ord_id:
                r_nak = self.client.get(f"/nakladnoy/{ord_id}")
                self.assertEqual(r_nak.status_code, 200)
                self.assertIn("Usta Dilshod", r_nak.data.decode("utf-8"))

                r_print = self.client.get(f"/nakladnoy/{ord_id}/print")
                self.assertEqual(r_print.status_code, 200)
                self.assertIn("Usta Dilshod", r_print.data.decode("utf-8"))

            # 4. Check /api/usta/<id>/orders
            r_orders = self.client.get(f"/api/usta/{usta1_id}/orders")
            self.assertEqual(r_orders.status_code, 200)
            d_orders = json.loads(r_orders.data)
            self.assertEqual(d_orders["summary"]["clients_count"], 2)
            self.assertEqual(d_orders["summary"]["orders_count"], 2)

            # 5. Check KPI dashboard for all requested periods: 1m, 3m, 6m, 1y, all
            for period in ["1m", "3m", "6m", "1y", "all"]:
                r_kpi = self.client.get(f"/kpi?period={period}")
                self.assertEqual(r_kpi.status_code, 200)
                html_kpi = r_kpi.data.decode("utf-8")
                self.assertIn("Ustalar & Prorablar Reytingi", html_kpi)
                self.assertIn("Usta Sherzod", html_kpi)
                self.assertIn("Usta Dilshod", html_kpi)

    def test_ustalar_under_xodimlar_section(self):
        # 1. Login as admin
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)

        # 2. Check Xodimlar page is clean and separate
        r_xod = self.client.get("/admin/xodimlar")
        self.assertEqual(r_xod.status_code, 200)
        html_xod = r_xod.data.decode("utf-8")
        self.assertIn("Xodimlar va Rollar Boshqaruvi", html_xod)
        self.assertIn("Ustalar", html_xod)  # in sidebar
        self.assertNotIn("Yangi Usta Qo'shish", html_xod)  # separate from xodimlar

        # 3. Check Ustalar page is standalone and dedicated
        r_ust = self.client.get("/admin/ustalar")
        self.assertEqual(r_ust.status_code, 200)
        html_ust = r_ust.data.decode("utf-8")
        self.assertIn("Ustalar va Prorablar Boshqaruvi", html_ust)
        self.assertIn("Yangi Usta Qo'shish", html_ust)
        self.assertIn("Jami Ustalar & Prorablar", html_ust)

        # 4. Add a new usta from /admin/ustalar/qoshish
        r_add = self.client.post("/admin/ustalar/qoshish", data={
            "full_name": "Hamid Plitochnik",
            "phone": "+998909876543",
            "customer_type": "usta",
            "notes": "Chilonzor brigadasi"
        }, follow_redirects=True)
        self.assertEqual(r_add.status_code, 200)
        html_after = r_add.data.decode("utf-8")
        self.assertIn("Hamid Plitochnik", html_after)
        self.assertIn("+998909876543", html_after)

        # 5. Delete the created usta
        conn = database.get_db()
        cur = conn.cursor()
        cur.execute("SELECT id FROM customers WHERE full_name = 'Hamid Plitochnik'")
        row = cur.fetchone()
        conn.close()
        self.assertIsNotNone(row)
        u_id = row["id"]

        r_del = self.client.post(f"/admin/ustalar/ochirish/{u_id}", follow_redirects=True)
        self.assertEqual(r_del.status_code, 200)
        self.assertNotIn("Hamid Plitochnik", r_del.data.decode("utf-8"))

    def test_comprehensive_full_system_verification(self):
        # 1. Login as admin
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)

        # 2. Check all main routes return 200 OK
        routes_to_test = [
            "/",
            "/ombor",
            "/ombor/yangi",
            "/ombor/siniqlar",
            "/sotuv",
            "/kassa",
            "/hisobot",
            "/admin/ustalar",
            "/admin/xodimlar",
            "/mijozlar",
            "/kpi",
            "/kpi?period=today",
            "/kpi?period=1m",
            "/kpi?period=3m",
            "/kpi?period=6m",
            "/kpi?period=1y",
            "/kpi?period=all",
        ]
        for route in routes_to_test:
            res = self.client.get(route)
            self.assertEqual(res.status_code, 200, f"Route {route} failed with status {res.status_code}")

        # 3. Check /hisobot elements
        r_hisobot = self.client.get("/hisobot")
        html_h = r_hisobot.data.decode("utf-8")
        
        # Verify 6 tabs present
        tabs = ["btn-generalTab", "btn-ordersTab", "btn-inboundTab", "btn-balanceTab", "btn-brandsTab", "btn-brokenTab"]
        for t in tabs:
            self.assertIn(t, html_h, f"Tab button {t} not found in /hisobot")

        tab_contents = ['id="generalTab"', 'id="ordersTab"', 'id="inboundTab"', 'id="balanceTab"', 'id="brandsTab"', 'id="brokenTab"']
        for tc in tab_contents:
            self.assertIn(tc, html_h, f"Tab content container {tc} not found in /hisobot")

        # Verify key headings and cards
        self.assertIn("Do'kon va Ombor Umumiy Jamlama Hisoboti", html_h)
        self.assertIn("Bosh Balans Vedomosti", html_h)
        self.assertIn("To'lov Usullari Bo'yicha Tushum Balansi", html_h)
        self.assertIn("Kafel Markalari Jamlama Aylanmasi & Ulushlari", html_h)
        self.assertIn("Hisobotni Chop Etish", html_h)

        # Verify NO edit or entry buttons in hisobot content
        self.assertNotIn('<a href="/ombor/tahrirlash', html_h)
        self.assertNotIn("Siniq Kiritish", html_h)
        self.assertNotIn("Yangi Savdo", html_h)

        # 4. Perform a real test sale with Usta and Payment method
        r_u = self.client.post("/api/ustalar/quick-add", json={"full_name": "Usta Botir", "phone": "+998931112233"})
        self.assertEqual(r_u.status_code, 200)
        u_data = json.loads(r_u.data)
        self.assertTrue(u_data.get("success"))
        botir_id = u_data.get("usta", {}).get("id")

        conn = database.get_db()
        cur = conn.cursor()
        cur.execute("SELECT * FROM products LIMIT 1")
        prod = cur.fetchone()
        conn.close()
        self.assertIsNotNone(prod)

        sale_payload = {
            "customer_name": "Qobil aka",
            "customer_phone": "+998901234567",
            "customer_type": "oddiy",
            "usta_id": botir_id,
            "payment_method": "karta",
            "items": [{
                "id": prod["id"],
                "brand": prod["brand"],
                "model_name": prod["model_name"],
                "price": prod["price"],
                "qty": 10.0
            }]
        }
        r_sale = self.client.post("/api/checkout", json=sale_payload)
        self.assertEqual(r_sale.status_code, 200)
        res_sale = json.loads(r_sale.data)
        self.assertTrue(res_sale.get("success"))
        order_id = res_sale.get("order_id")

        # 5. Check order documents
        r_nak = self.client.get(f"/nakladnoy/{order_id}")
        self.assertEqual(r_nak.status_code, 200)
        self.assertIn("Qobil aka", r_nak.data.decode("utf-8"))
        self.assertIn("Usta Botir", r_nak.data.decode("utf-8"))

        r_chek = self.client.get(f"/nakladnoy/{order_id}/chek")
        self.assertEqual(r_chek.status_code, 200)

        r_pr = self.client.get(f"/nakladnoy/{order_id}/print")
        self.assertEqual(r_pr.status_code, 200)

        # 6. Dispatch order from warehouse
        r_disp = self.client.post(f"/api/order/{order_id}/dispatch", json={"note": "Yuk to'liq topshirildi"})
        self.assertEqual(r_disp.status_code, 200)

        # 7. Check Hisobot updated values
        r_hisobot2 = self.client.get("/hisobot")
        html_h2 = r_hisobot2.data.decode("utf-8")
        self.assertIn("Qobil aka", html_h2)
        self.assertIn("10.00", html_h2)

        # 8. Check KPI dashboard updated
        r_kpi2 = self.client.get("/kpi?period=1m")
        html_kpi2 = r_kpi2.data.decode("utf-8")
        self.assertIn("Usta Botir", html_kpi2)

class TestCyberSecurityAndResponsiveness(unittest.TestCase):
    def setUp(self):
        self.client = app.app.test_client()
        app.app.config["TESTING"] = True
        app.reset_all_rate_limits()

    def tearDown(self):
        app.reset_all_rate_limits()

    def test_security_headers(self):
        """Kiberxavfsizlik: HTTP xavfsizlik sarlavhalari (Anti-Clickjacking, Anti-XSS, MIME sniffing)"""
        r = self.client.get("/login")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers.get("X-Content-Type-Options"), "nosniff")
        self.assertEqual(r.headers.get("X-Frame-Options"), "SAMEORIGIN")
        self.assertEqual(r.headers.get("X-XSS-Protection"), "1; mode=block")
        self.assertEqual(r.headers.get("Referrer-Policy"), "strict-origin-when-cross-origin")
        self.assertIn("geolocation=()", r.headers.get("Permissions-Policy", ""))

        # Authenticated user should have Cache-Control: no-store
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r_auth = self.client.get("/")
        self.assertIn("no-store", r_auth.headers.get("Cache-Control", ""))

    def test_brute_force_lockout_and_rate_limiting(self):
        """Buzib kirishdan himoya (Anti-Brute-Force): 5 ta xato urinishdan so'ng 429 bloklash"""
        app.reset_all_rate_limits()

        # 1-4 ta xato urinish (ogohlantirish va qolgan urinishlar ko'rsatiladi)
        for i in range(4):
            r = self.client.post("/login", data={"username": "admin", "password": f"wrong_pwd_{i}"})
            self.assertEqual(r.status_code, 200, f"Attempt {i+1} should return 200 login page")
            self.assertIn(b"Login yoki parol", r.data)

        # 5-urinish: Oxirgi urinish xato bo'lgach, bloklash xabari chiqadi
        r5 = self.client.post("/login", data={"username": "admin", "password": "wrong_pwd_5"})
        self.assertEqual(r5.status_code, 200)
        self.assertIn(b"15 daqiqaga bloklandi", r5.data)

        # 6-urinish: Tizim to'liq bloklangan, HTTP 429 qaytaradi
        r6 = self.client.post("/login", data={"username": "admin", "password": "wrong_pwd_6"})
        self.assertEqual(r6.status_code, 429, "Attempt after 5 failures must trigger HTTP 429 rate limit")
        self.assertTrue(b"bloklandi" in r6.data or b"Buzib kirishdan himoya" in r6.data)

        # 7-urinish: Hatto to'g'ri parol kiritilganda ham blok holatida qoladi
        r7 = self.client.post("/login", data={"username": "admin", "password": "admin123"})
        self.assertEqual(r7.status_code, 429)

        # Rate limit tozalanadi (masalan, 15 daqiqadan so'ng)
        app.reset_all_rate_limits()
        self.client.get("/logout")
        r_ok = self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        self.assertEqual(r_ok.status_code, 200)
        self.assertIn(b"Boshqaruv Paneli", r_ok.data)

    def test_open_redirect_protection(self):
        """Kiberxavfsizlik: Open Redirect phishing hujumlaridan himoya"""
        app.reset_all_rate_limits()
        self.client.get("/logout")

        # Tashqi firibgar saytga yo'naltirish urinishi
        r_phish = self.client.post("/login?next=https://evil-phishing-attacker.com/steal", data={
            "username": "admin",
            "password": "admin123"
        }, follow_redirects=False)

        # Tashqi saytga redirect qilmasligi kerak
        redirect_loc = r_phish.headers.get("Location", "")
        self.assertNotIn("evil-phishing-attacker.com", redirect_loc)

        # Sessiyani tozalab, tizim ichki xavfsiz sahifaga yo'naltirishini tekshirish
        self.client.get("/logout")
        r_safe = self.client.post("/login?next=/ombor", data={
            "username": "admin",
            "password": "admin123"
        }, follow_redirects=False)
        self.assertEqual(r_safe.status_code, 302)
        self.assertIn("/ombor", r_safe.headers.get("Location", ""))

    def test_cryptographic_secret_and_session_hardening(self):
        """Kriptografik 256-bit Secret Key va sessiya cookie himoyasi"""
        self.assertTrue(len(app.app.secret_key) >= 32)
        self.assertTrue(app.app.config.get("SESSION_COOKIE_HTTPONLY"))
        self.assertEqual(app.app.config.get("SESSION_COOKIE_SAMESITE"), "Lax")
        self.assertEqual(app.app.config.get("SESSION_COOKIE_NAME"), "kafel_secure_session")
        self.assertEqual(app.app.config.get("PERMANENT_SESSION_LIFETIME").total_seconds(), 28800)

    def test_responsive_viewport_and_mobile_classes(self):
        """Mobil telefon, planshet va noutbuk moslashuvchanligi (Viewport va CSS)"""
        # Login sahifasida viewport mavjudligi
        r_login = self.client.get("/login")
        self.assertEqual(r_login.status_code, 200)
        self.assertIn('name="viewport"', r_login.data.decode("utf-8"))

        # Base template viewport va hisobot sahifasi moslashuvchanligi
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r_hisobot = self.client.get("/hisobot")
        html_h = r_hisobot.data.decode("utf-8")
        self.assertIn('name="viewport"', html_h)
        self.assertIn('content="width=device-width, initial-scale=1.0', html_h)
        # 7 tab tugmalari va responsive classlar
        self.assertIn('id="tabsNav"', html_h)
        self.assertIn('grid-cols-2', html_h)
        self.assertIn('overflow-x-auto', html_h)

    def test_action_logs_audit_trail(self):
        """Amallar audit jurnali (kiritdi, tahrirladi, ochirdi, sotdi, topshirdi) tekshiruvi"""
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r_fact = self.client.post("/api/factories/add", json={"name": "Audit Test Zavod", "country": "Italiya"})
        self.assertEqual(r_fact.status_code, 200)

        r_usta = self.client.post("/api/ustalar/quick-add", json={"full_name": "Audit Test Usta", "phone": "+998901112233"})
        self.assertEqual(r_usta.status_code, 200)

        logs = database.get_action_logs(limit=50)
        action_types = [l["action_type"] for l in logs]
        self.assertIn("kiritdi", action_types)

        kiritdi_logs = database.get_action_logs(action_type="kiritdi")
        self.assertTrue(len(kiritdi_logs) > 0)
        self.assertTrue(all(l["action_type"] == "kiritdi" for l in kiritdi_logs))

    def test_seller_sold_products_isolation_and_admin_view(self):
        """Sotuvchining o'z savdolari va tovarlar ro'yxati alohidaligi hamda admin ko'rinishi"""
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        self.client.post("/admin/xodimlar", data={
            "full_name": "Sardor Sotuvchi",
            "username": "sardor_seller",
            "password": "sardorpassword123",
            "role": "sotuvchi"
        }, follow_redirects=True)

        self.client.get("/logout")
        self.client.post("/login", data={"username": "sardor_seller", "password": "sardorpassword123"}, follow_redirects=True)

        conn = database.get_db()
        cursor = conn.cursor()
        cursor.execute("SELECT id, price, quantity_in_stock FROM products WHERE quantity_in_stock >= 5 LIMIT 1")
        p = cursor.fetchone()
        conn.close()
        self.assertIsNotNone(p)

        r_pay = self.client.post("/api/checkout", json={
            "items": [{"id": p["id"], "qty": 2.0, "price": p["price"]}],
            "customer_name": "Sardor Xaridori",
            "payment_method": "naqd"
        })
        self.assertEqual(r_pay.status_code, 200)

        r_hisobot_sardor = self.client.get("/hisobot")
        self.assertEqual(r_hisobot_sardor.status_code, 200)
        html_sardor = r_hisobot_sardor.data.decode("utf-8")
        self.assertIn("Mening Sotgan Kafellarim Ro'yxati", html_sardor)
        self.assertIn("Sardor Xaridori", html_sardor)

        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r_hisobot_admin = self.client.get("/hisobot")
        self.assertEqual(r_hisobot_admin.status_code, 200)
        html_admin = r_hisobot_admin.data.decode("utf-8")
        self.assertIn("Har Bir Sotuvchining Sotgan Tovarlari Spiskasi", html_admin)
        self.assertIn("Sardor Sotuvchi", html_admin)
        self.assertIn("Amallar Tarixi", html_admin)

    def test_hisobot_role_separation(self):
        """Hisobot sahifasining har bir rol (sotuvchi, omborchi, admin) bo'yicha to'liq ajratilganligini tekshirish"""
        # 1. Sotuvchi ko'rinishi
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"}, follow_redirects=True)
        r_sotuv = self.client.get("/hisobot")
        self.assertEqual(r_sotuv.status_code, 200)
        html_sotuv = r_sotuv.data.decode("utf-8")

        # Sotuvchi o'ziga tegishli sarlavha va tablarni ko'radi
        self.assertIn("Mening Shaxsiy Savdo Hisobotim", html_sotuv)
        self.assertIn('id="ordersTab"', html_sotuv)
        self.assertIn('id="logsTab"', html_sotuv)
        self.assertIn("Mening Amallarim Tarixi (Shaxsiy Audit)", html_sotuv)

        # Sotuvchiga yopiq bo'lgan tablar mutlaqo render qilinmasligi kerak
        self.assertNotIn('id="generalTab"', html_sotuv)
        self.assertNotIn('id="inboundTab"', html_sotuv)
        self.assertNotIn('id="balanceTab"', html_sotuv)
        self.assertNotIn('id="brandsTab"', html_sotuv)
        self.assertNotIn('id="brokenTab"', html_sotuv)

        # 2. Omborchi ko'rinishi
        self.client.get("/logout")
        self.client.post("/login", data={"username": "omborchi", "password": "omborchi123"}, follow_redirects=True)
        r_ombor = self.client.get("/hisobot")
        self.assertEqual(r_ombor.status_code, 200)
        html_ombor = r_ombor.data.decode("utf-8")

        # Omborchi o'ziga tegishli sarlavha va tablarni ko'radi
        self.assertIn("Omborxona & Logistika Hisoboti", html_ombor)
        self.assertIn('id="inboundTab"', html_ombor)
        self.assertIn('id="balanceTab"', html_ombor)
        self.assertIn('id="ordersTab"', html_ombor)
        self.assertIn('id="brandsTab"', html_ombor)
        self.assertIn('id="brokenTab"', html_ombor)
        self.assertIn('id="logsTab"', html_ombor)
        self.assertIn("Ombor Amallari Tarixi (Logistika Auditi)", html_ombor)

        # Omborchiga moliyaviy umumiy hisobot va pul ustunlari ko'rinmasligi kerak
        self.assertNotIn('id="generalTab"', html_ombor)
        self.assertNotIn("Zaxira Qiymati", html_ombor)
        self.assertNotIn("Kvadrat Narxi", html_ombor)
        self.assertNotIn("Yetkazilgan Zarar", html_ombor)

        # 3. Admin ko'rinishi
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)
        r_admin = self.client.get("/hisobot")
        self.assertEqual(r_admin.status_code, 200)
        html_admin = r_admin.data.decode("utf-8")

        # Admin barcha 7 ta tabni va to'liq moliyaviy hisobotni ko'radi
        self.assertIn("Ombor & Savdo Audit Jurnali", html_admin)
        self.assertIn('id="generalTab"', html_admin)
        self.assertIn('id="ordersTab"', html_admin)
        self.assertIn('id="inboundTab"', html_admin)
        self.assertIn('id="balanceTab"', html_admin)
        self.assertIn('id="brandsTab"', html_admin)
        self.assertIn('id="brokenTab"', html_admin)
        self.assertIn('id="logsTab"', html_admin)
        self.assertIn("Zaxira Qiymati", html_admin)
        self.assertIn("Yetkazilgan Zarar", html_admin)
        self.assertIn("Tizimdagi Barcha Amallar Tarixi (To'liq Audit)", html_admin)

    def test_excel_exports_all_endpoints(self):
        """Barcha 12 ta alohida Excel eksport marshrutlari to'g'ri ishlashi va haqiqiy .xlsx fayl qaytarishi"""
        import io
        import openpyxl

        # 1. Admin sifatida tizimga kirish
        self.client.get("/logout")
        self.client.post("/login", data={"username": "admin", "password": "admin123"}, follow_redirects=True)

        endpoints = [
            ("/export/excel/ombor", "Ombor Qoldiqlari"),
            ("/export/excel/hisobot/savdolar", "Savdolar Hisoboti"),
            ("/export/excel/hisobot/kirim", "Kirim Tarixi"),
            ("/export/excel/hisobot/balans", "Qoldiq & Balans"),
            ("/export/excel/hisobot/markalar", "Kafel Markalari"),
            ("/export/excel/hisobot/siniqlar", "Siniq & Braklar"),
            ("/export/excel/hisobot/audit-log", "Amallar Tarixi"),
            ("/export/excel/hisobot/umumiy", "Bosh Balans"),
            ("/export/excel/nakladnoylar", "Nakladnoylar"),
            ("/export/excel/mijozlar", "Mijozlar & Nasiya"),
            ("/export/excel/xarajatlar", "Xarajatlar"),
            ("/export/excel/ustalar", "Ustalar & Prorablar")
        ]

        for url, expected_sheet_title in endpoints:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, f"{url} status 200 bo'lishi kerak")
            self.assertTrue(
                r.headers["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                f"{url} Content-Type Excel bo'lishi kerak, lekin {r.headers['Content-Type']} chiqdi"
            )
            
            # Excel fayli haqiqiy ekanligini openpyxl orqali tekshirish
            wb = openpyxl.load_workbook(io.BytesIO(r.data))
            sheet = wb.active
            self.assertEqual(sheet.cell(row=1, column=1).value, "KAFEL CENTER & SAVDO TIZIMI")
            self.assertTrue(len(sheet.title) > 0)
            # 5-qator sarlavhalar mavjudligini tekshirish
            self.assertIsNotNone(sheet.cell(row=5, column=1).value)

    def test_excel_exports_role_isolation(self):
        """Excel eksportida rollar bo'yicha ruxsatlar va pul ko'rsatkichlarini yashirish (omborchi/sotuvchi)"""
        import io
        import openpyxl

        # 1. Omborchi ko'rinishi: ombor qoldig'i Excelida narxlar bo'lmasligi kerak
        self.client.get("/logout")
        self.client.post("/login", data={"username": "omborchi", "password": "omborchi123"}, follow_redirects=True)

        r_ombor = self.client.get("/export/excel/ombor")
        self.assertEqual(r_ombor.status_code, 200)
        wb_ombor = openpyxl.load_workbook(io.BytesIO(r_ombor.data))
        ws_ombor = wb_ombor.active
        header_vals_ombor = [ws_ombor.cell(row=5, column=c).value for c in range(1, ws_ombor.max_column + 1)]
        self.assertNotIn("1 m² Narxi", header_vals_ombor)
        self.assertNotIn("Zaxira Qiymati", header_vals_ombor)

        # Omborchi umumiy moliyaviy balansga kira olmasligi kerak (redirect)
        r_umumiy = self.client.get("/export/excel/hisobot/umumiy", follow_redirects=False)
        self.assertEqual(r_umumiy.status_code, 302)

        # 2. Sotuvchi ko'rinishi: sotuvchi ombor qoldig'ini eksport qila olmasligi kerak
        self.client.get("/logout")
        self.client.post("/login", data={"username": "sotuvchi", "password": "sotuvchi123"}, follow_redirects=True)

        r_sotuv_ombor = self.client.get("/export/excel/ombor", follow_redirects=False)
        self.assertEqual(r_sotuv_ombor.status_code, 302)

        # Sotuvchi o'z savdolarini eksport qila oladi
        r_sotuv_savdo = self.client.get("/export/excel/hisobot/savdolar")
        self.assertEqual(r_sotuv_savdo.status_code, 200)

if __name__ == "__main__":
    unittest.main()



